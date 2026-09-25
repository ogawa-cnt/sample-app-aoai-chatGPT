import io
import json
import re


def _parse_markdown_blocks(markdown_text: str) -> list:
    # 見出し(#/##/###)・箇条書き(-)・番号付きリスト(1.)・表(| | |)・段落、という
    # 限定的なMarkdownの書式だけを対象にした簡易パーサー
    lines = markdown_text.splitlines()
    blocks = []
    i = 0
    n = len(lines)

    while i < n:
        stripped = lines[i].strip()
        if not stripped:
            i += 1
            continue

        heading_match = re.match(r"^(#{1,3})\s+(.*)$", stripped)
        if heading_match:
            blocks.append({
                "type": "heading",
                "level": len(heading_match.group(1)),
                "text": heading_match.group(2).strip(),
            })
            i += 1
            continue

        if re.match(r"^[-*]\s+", stripped):
            items = []
            while i < n and re.match(r"^[-*]\s+", lines[i].strip()):
                items.append(re.sub(r"^[-*]\s+", "", lines[i].strip()))
                i += 1
            blocks.append({"type": "bullet_list", "items": items})
            continue

        if re.match(r"^\d+\.\s+", stripped):
            items = []
            while i < n and re.match(r"^\d+\.\s+", lines[i].strip()):
                items.append(re.sub(r"^\d+\.\s+", "", lines[i].strip()))
                i += 1
            blocks.append({"type": "numbered_list", "items": items})
            continue

        if stripped.startswith("|"):
            table_lines = []
            while i < n and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            rows = [
                [cell.strip() for cell in line.strip("|").split("|")]
                for line in table_lines
            ]
            if len(rows) >= 2 and all(re.match(r"^:?-+:?$", c) for c in rows[1]):
                header, data_rows = rows[0], rows[2:]
            else:
                header, data_rows = rows[0], rows[1:]
            blocks.append({"type": "table", "header": header, "rows": data_rows})
            continue

        para_lines = [stripped]
        i += 1
        while i < n and lines[i].strip() and not re.match(r"^(#{1,3}\s+|[-*]\s+|\d+\.\s+|\|)", lines[i].strip()):
            para_lines.append(lines[i].strip())
            i += 1
        blocks.append({"type": "paragraph", "text": " ".join(para_lines)})

    return blocks


def _generate_docx(markdown_text: str) -> bytes:
    from docx import Document

    document = Document()
    for block in _parse_markdown_blocks(markdown_text):
        if block["type"] == "heading":
            document.add_heading(block["text"], level=block["level"])
        elif block["type"] == "paragraph":
            document.add_paragraph(block["text"])
        elif block["type"] == "bullet_list":
            for item in block["items"]:
                document.add_paragraph(item, style="List Bullet")
        elif block["type"] == "numbered_list":
            for item in block["items"]:
                document.add_paragraph(item, style="List Number")
        elif block["type"] == "table":
            header = block["header"]
            table = document.add_table(rows=1, cols=len(header))
            table.style = "Light Grid Accent 1"
            for idx, cell_text in enumerate(header):
                table.rows[0].cells[idx].text = cell_text
            for row in block["rows"]:
                cells = table.add_row().cells
                for idx, cell_text in enumerate(row):
                    if idx < len(cells):
                        cells[idx].text = cell_text

    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _add_tf_paragraph(text_frame, text: str, bold: bool = False):
    if len(text_frame.paragraphs) == 1 and not text_frame.paragraphs[0].runs and not text_frame.paragraphs[0].text:
        paragraph = text_frame.paragraphs[0]
    else:
        paragraph = text_frame.add_paragraph()
    paragraph.text = text
    if bold:
        for run in paragraph.runs:
            run.font.bold = True
    return paragraph


def _generate_pptx(markdown_text: str) -> bytes:
    from pptx import Presentation

    presentation = Presentation()
    slide = None
    body_tf = None

    def ensure_slide():
        nonlocal slide, body_tf
        if slide is None:
            layout = presentation.slide_layouts[1]
            slide = presentation.slides.add_slide(layout)
            slide.shapes.title.text = ""
            body_tf = slide.placeholders[1].text_frame
            body_tf.clear()

    for block in _parse_markdown_blocks(markdown_text):
        # 見出し(#/##)ごとに新しいスライドを作る
        if block["type"] == "heading" and block["level"] <= 2:
            layout = presentation.slide_layouts[1]
            slide = presentation.slides.add_slide(layout)
            slide.shapes.title.text = block["text"]
            body_tf = slide.placeholders[1].text_frame
            body_tf.clear()
            continue

        ensure_slide()

        if block["type"] == "heading":
            _add_tf_paragraph(body_tf, block["text"], bold=True)
        elif block["type"] == "paragraph":
            _add_tf_paragraph(body_tf, block["text"])
        elif block["type"] in ("bullet_list", "numbered_list"):
            for item in block["items"]:
                _add_tf_paragraph(body_tf, item)
        elif block["type"] == "table":
            for row in [block["header"]] + block["rows"]:
                _add_tf_paragraph(body_tf, " | ".join(row))

    if slide is None:
        presentation.slides.add_slide(presentation.slide_layouts[1])

    buf = io.BytesIO()
    presentation.save(buf)
    return buf.getvalue()


def _sanitize_sheet_title(title: str) -> str:
    title = re.sub(r'[:\\/?*\[\]]', "_", title).strip()
    return title[:31] or "Sheet1"


def _generate_xlsx(markdown_text: str) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    current_sheet = None

    for block in _parse_markdown_blocks(markdown_text):
        if block["type"] == "heading":
            current_sheet = workbook.create_sheet(title=_sanitize_sheet_title(block["text"]))
            continue

        if current_sheet is None:
            current_sheet = workbook.create_sheet(title="Sheet1")

        if block["type"] == "table":
            current_sheet.append(block["header"])
            for row in block["rows"]:
                current_sheet.append(row)
        elif block["type"] in ("bullet_list", "numbered_list"):
            for item in block["items"]:
                current_sheet.append([item])
        elif block["type"] == "paragraph":
            current_sheet.append([block["text"]])

    if not workbook.sheetnames:
        workbook.create_sheet(title="Sheet1")

    buf = io.BytesIO()
    workbook.save(buf)
    return buf.getvalue()


def generate_file(ext: str, content: str) -> bytes:
    ext = ext.lower()

    if ext in ("txt", "md"):
        return content.encode("utf-8")

    if ext == "json":
        parsed = json.loads(content)
        return json.dumps(parsed, ensure_ascii=False, indent=2).encode("utf-8")

    if ext == "html":
        return content.encode("utf-8")

    if ext == "docx":
        return _generate_docx(content)

    if ext == "xlsx":
        return _generate_xlsx(content)

    if ext == "pptx":
        return _generate_pptx(content)

    raise ValueError(f"Unsupported generation format: .{ext}")
