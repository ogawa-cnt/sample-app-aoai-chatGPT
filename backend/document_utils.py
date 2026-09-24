import io
import re


def _strip_html_tags(html_text: str) -> str:
    # <script>と<style>の中身は読み取っても意味がないので丸ごと除去
    html_text = re.sub(r"<(script|style)[^>]*>.*?</\1>", "", html_text, flags=re.DOTALL | re.IGNORECASE)
    # 残りのタグを除去
    text = re.sub(r"<[^>]+>", " ", html_text)
    # 連続する空白をまとめる
    return re.sub(r"[ \t]+", " ", text).strip()


def _iter_docx_block_items(document):
    # 段落と表を、文書内に登場する順番のまま取り出すための定石の実装
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield Table(child, document)


def _extract_text_from_docx(file_bytes: bytes) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(io.BytesIO(file_bytes))
    parts = []
    for block in _iter_docx_block_items(document):
        if isinstance(block, Paragraph):
            if block.text.strip():
                parts.append(block.text)
        elif isinstance(block, Table):
            for row in block.rows:
                cells = [cell.text.strip() for cell in row.cells]
                parts.append("| " + " | ".join(cells) + " |")
    return "\n".join(parts)


MAX_ROWS_PER_SHEET = 2000


def _extract_text_from_xlsx(file_bytes: bytes) -> str:
    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    sheet_parts = []
    for sheet in workbook.worksheets:
        rows = [
            row for row in sheet.iter_rows(values_only=True)
            if any(cell is not None for cell in row)
        ]
        if not rows:
            continue

        header, *data_rows = rows
        header_cells = ["" if v is None else str(v) for v in header]

        lines = [f"## シート: {sheet.title}"]
        lines.append("| " + " | ".join(header_cells) + " |")
        lines.append("| " + " | ".join(["---"] * len(header_cells)) + " |")
        for row in data_rows[:MAX_ROWS_PER_SHEET]:
            cells = ["" if v is None else str(v) for v in row]
            lines.append("| " + " | ".join(cells) + " |")
        if len(data_rows) > MAX_ROWS_PER_SHEET:
            lines.append(f"...(以下省略、{len(data_rows) - MAX_ROWS_PER_SHEET}行省略)")

        sheet_parts.append("\n".join(lines))

    return "\n\n".join(sheet_parts)


def _extract_text_from_pptx(file_bytes: bytes) -> str:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(file_bytes))
    slide_parts = []
    for index, slide in enumerate(presentation.slides, start=1):
        lines = [f"## スライド {index}"]
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = "\n".join(
                    p.text for p in shape.text_frame.paragraphs if p.text.strip()
                )
                if text:
                    lines.append(text)
            if shape.has_table:
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    lines.append("| " + " | ".join(cells) + " |")

        if slide.has_notes_slide:
            notes_text = slide.notes_slide.notes_text_frame.text.strip()
            if notes_text:
                lines.append(f"(発表者ノート) {notes_text}")

        slide_parts.append("\n".join(lines))

    return "\n\n".join(slide_parts)


def extract_text_from_file(file_bytes: bytes, filename: str) -> str:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if ext in ("txt", "md", "json"):
        return file_bytes.decode("utf-8", errors="ignore")

    if ext in ("html", "htm"):
        raw_html = file_bytes.decode("utf-8", errors="ignore")
        return _strip_html_tags(raw_html)

    if ext == "pdf":
        import pdfplumber
        text_parts = []
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
        return "\n".join(text_parts)

    if ext == "docx":
        return _extract_text_from_docx(file_bytes)

    if ext == "xlsx":
        return _extract_text_from_xlsx(file_bytes)

    if ext == "pptx":
        return _extract_text_from_pptx(file_bytes)

    raise ValueError(f"Unsupported file type: .{ext}")


MAX_EXTRACTED_CHARS = 200000


def truncate_text(text: str, max_chars: int = MAX_EXTRACTED_CHARS) -> str:
    if len(text) <= max_chars:
        return text
    return text[:max_chars] + "\n\n...(以下省略、文字数上限のため切り捨てられました)"
